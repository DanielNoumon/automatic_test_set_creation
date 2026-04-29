"""
Word (.docx) parser: extracts structured sections from Word documents.

Uses python-docx to read paragraphs and their styles, mapping
Word heading styles (Heading 1, Heading 2, …) to Section levels.
Falls back to paragraph-based splitting when no headings are found.
"""
import re
from typing import List
from pathlib import Path

from docx import Document as DocxDocument

from .document import Document, Section
from .pdf_parser import _tag_section


# Heading style pattern: "Heading 1", "Heading 2", etc.
_HEADING_STYLE_RE = re.compile(r"^Heading\s+(\d+)$", re.IGNORECASE)


def parse_docx(filepath: str) -> Document:
    """Parse a .docx file into a structured Document."""
    path = Path(filepath)
    doc = DocxDocument(str(path))

    # Extract all text for raw_text
    all_paragraphs = []
    for para in doc.paragraphs:
        all_paragraphs.append(para.text)
    raw_text = "\n".join(all_paragraphs)

    # Try heading-based splitting first
    sections = _parse_with_headings(doc)
    if not sections:
        sections = _parse_paragraphs(raw_text)

    for sec in _flatten(sections):
        _tag_section(sec)

    return Document(
        filename=path.name,
        raw_text=raw_text,
        sections=sections,
        page_count=1,
        doc_type="docx",
    )


def _parse_with_headings(doc: DocxDocument) -> List[Section]:
    """Split by Word heading styles."""
    entries = []  # (level, heading_text, content_lines)
    current_heading = None
    current_level = 0
    current_lines = []

    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""
        m = _HEADING_STYLE_RE.match(style_name)

        if m:
            # Save previous section
            if current_heading is not None:
                entries.append((
                    current_level,
                    current_heading,
                    "\n".join(current_lines).strip(),
                ))
            current_level = int(m.group(1)) - 1
            current_heading = para.text.strip()
            current_lines = []
        else:
            text = para.text.strip()
            if text:
                current_lines.append(text)

    # Save last section
    if current_heading is not None:
        entries.append((
            current_level,
            current_heading,
            "\n".join(current_lines).strip(),
        ))

    if not entries:
        return []

    sections = []
    pos = 0
    for level, heading, content in entries:
        sec = Section(
            heading=heading,
            content=content,
            level=level,
            page_start=1,
            page_end=1,
            char_start=pos,
            char_end=pos + len(heading) + len(content),
        )
        pos += len(heading) + len(content) + 2
        sections.append(sec)

    return sections


def _parse_paragraphs(raw_text: str) -> List[Section]:
    """Fallback: split into paragraph-based sections."""
    paragraphs = re.split(r"\n\s*\n", raw_text)
    sections = []
    pos = 0
    for para in paragraphs:
        para = para.strip()
        if not para or len(para) < 20:
            pos += len(para) + 2
            continue
        lines = para.split("\n")
        heading = lines[0].strip()[:80]
        content = (
            "\n".join(lines[1:]).strip()
            if len(lines) > 1 else para
        )
        sections.append(Section(
            heading=heading,
            content=content,
            level=0,
            page_start=1,
            page_end=1,
            char_start=pos,
            char_end=pos + len(para),
        ))
        pos += len(para) + 2
    return sections


def _flatten(sections: List[Section]) -> List[Section]:
    result = []
    for sec in sections:
        result.append(sec)
        result.extend(_flatten(sec.children))
    return result
